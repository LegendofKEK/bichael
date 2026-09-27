"""Generate missing early craft / bait / base-mat icons (Bellgrave item inventory style)."""
from __future__ import annotations

import math
from pathlib import Path

from openpyxl import load_workbook
from PIL import Image, ImageDraw

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "apps" / "web" / "public" / "icons" / "items"
MASTERS = ROOT / "assets" / "icons" / "items"
MANIFEST = ROOT / "docs" / "legend-of-kek-item-manifest.xlsx"
SIZE = 96
MAGENTA = (255, 0, 255, 255)


def new_canvas():
    im = Image.new("RGBA", (SIZE, SIZE), MAGENTA)
    return im, ImageDraw.Draw(im)


def save(im: Image.Image, name: str) -> None:
    px = im.load()
    assert px is not None
    for y in range(SIZE):
        for x in range(SIZE):
            r, g, b, _a = px[x, y]
            if r > 240 and g < 20 and b > 240:
                px[x, y] = (0, 0, 0, 0)
    OUT.mkdir(parents=True, exist_ok=True)
    MASTERS.mkdir(parents=True, exist_ok=True)
    path = OUT / name
    im.save(path, optimize=True)
    im.save(MASTERS / name, optimize=True)
    print("wrote", name)


def bone_ring(im, d):
    # ivory ring with small bone knuckle
    d.ellipse([22, 22, 74, 74], outline=(230, 220, 200, 255), width=10)
    d.ellipse([30, 30, 66, 66], outline=(190, 175, 150, 255), width=3)
    d.ellipse([40, 16, 56, 34], fill=(235, 225, 205, 255), outline=(170, 150, 120, 255))
    d.line([(44, 20), (52, 30)], fill=(160, 140, 110, 255), width=1)


def glass_bead(im, d):
    d.ellipse([28, 28, 68, 68], fill=(140, 210, 220, 255), outline=(80, 140, 160, 255))
    d.ellipse([36, 34, 52, 50], fill=(220, 245, 250, 255))
    d.ellipse([44, 48, 58, 60], fill=(100, 170, 190, 180))
    d.ellipse([46, 22, 50, 28], fill=(230, 240, 245, 255))  # string hole hint


def distilled_ash(im, d):
    # small vial of pale ash liquid
    d.polygon([(36, 28), (60, 28), (64, 78), (32, 78)], fill=(200, 210, 215, 255), outline=(120, 130, 140, 255))
    d.rectangle([40, 18, 56, 30], fill=(160, 150, 130, 255), outline=(100, 90, 70, 255))
    d.polygon([(38, 48), (58, 48), (60, 74), (36, 74)], fill=(210, 205, 190, 255))
    d.ellipse([42, 52, 50, 60], fill=(240, 235, 220, 200))


def dried_ration(im, d):
    # wrapped ration brick
    d.rounded_rectangle([22, 34, 74, 70], radius=4, fill=(160, 120, 70, 255), outline=(100, 70, 40, 255))
    d.rectangle([28, 40, 68, 64], fill=(190, 150, 95, 255))
    d.line([(34, 48), (62, 48)], fill=(120, 85, 50, 255), width=2)
    d.line([(34, 56), (62, 56)], fill=(120, 85, 50, 255), width=2)
    d.ellipse([44, 28, 52, 36], fill=(140, 100, 55, 255))  # tie


def pale_bait(im, d):
    # pale grub / worm on hook
    d.ellipse([30, 40, 70, 68], fill=(230, 220, 190, 255), outline=(170, 155, 120, 255))
    d.ellipse([50, 34, 72, 56], fill=(240, 230, 200, 255), outline=(180, 160, 125, 255))
    d.arc([20, 18, 48, 50], 200, 40, fill=(160, 160, 170, 255), width=3)
    d.line([(24, 22), (20, 14)], fill=(140, 140, 150, 255), width=2)


def spice_rub(im, d):
    # open pouch of reddish spice
    d.polygon([(28, 70), (20, 42), (48, 28), (76, 42), (68, 70)], fill=(140, 90, 50, 255), outline=(90, 55, 30, 255))
    d.ellipse([30, 38, 66, 62], fill=(190, 80, 45, 255), outline=(130, 50, 30, 255))
    for x, y in [(38, 46), (48, 50), (56, 44), (44, 54)]:
        d.ellipse([x - 2, y - 2, x + 2, y + 2], fill=(220, 140, 70, 255))


def smoked_minnow(im, d):
    d.ellipse([18, 40, 70, 64], fill=(120, 90, 60, 255), outline=(70, 50, 35, 255))
    d.polygon([(70, 52), (86, 42), (86, 62)], fill=(100, 75, 50, 255))
    d.ellipse([26, 46, 34, 54], fill=(40, 30, 20, 255))
    d.line([(36, 52), (62, 52)], fill=(80, 60, 40, 255), width=1)
    # smoke wisps
    d.arc([40, 18, 58, 40], 200, 320, fill=(180, 180, 180, 180), width=2)
    d.arc([50, 14, 68, 36], 200, 320, fill=(200, 200, 200, 160), width=2)


def ash_dust(im, d):
    d.ellipse([26, 40, 70, 78], fill=(150, 145, 140, 255), outline=(90, 88, 85, 255))
    d.ellipse([34, 34, 62, 58], fill=(180, 175, 170, 255))
    d.ellipse([40, 30, 56, 44], fill=(210, 205, 200, 255))


def bone_chip(im, d):
    d.polygon([(30, 28), (58, 22), (70, 50), (52, 78), (26, 60)], fill=(235, 225, 205, 255), outline=(170, 155, 130, 255))
    d.line([(40, 36), (55, 55)], fill=(190, 175, 150, 255), width=2)
    d.ellipse([44, 42, 52, 50], fill=(210, 200, 180, 255))


def dried_ration_scrap(im, d):
    d.polygon([(24, 36), (60, 28), (76, 58), (40, 72)], fill=(170, 130, 80, 255), outline=(110, 75, 45, 255))
    d.line([(34, 44), (58, 52)], fill=(130, 95, 55, 255), width=2)
    d.line([(38, 56), (62, 48)], fill=(130, 95, 55, 255), width=1)


def desiccated_hide(im, d):
    d.polygon([(22, 50), (34, 24), (70, 28), (78, 58), (58, 78), (28, 72)], fill=(150, 105, 70, 255), outline=(95, 60, 40, 255))
    d.arc([32, 36, 66, 64], 20, 200, fill=(120, 80, 50, 255), width=2)
    d.ellipse([44, 40, 54, 50], fill=(130, 90, 55, 255))


def antidote(im, d):
    d.ellipse([34, 22, 62, 42], fill=(90, 150, 80, 255), outline=(50, 100, 50, 255))
    d.polygon([(38, 40), (58, 40), (64, 78), (32, 78)], fill=(120, 190, 110, 255), outline=(60, 120, 60, 255))
    d.rectangle([42, 16, 54, 26], fill=(140, 120, 80, 255))


def ashbeam_cane(im, d):
    d.line([(48, 18), (48, 82)], fill=(180, 150, 90, 255), width=8)
    d.ellipse([38, 12, 58, 28], fill=(200, 170, 100, 255), outline=(120, 90, 50, 255))
    d.ellipse([44, 72, 52, 84], fill=(100, 80, 50, 255))


def bone_earring(im, d):
    d.ellipse([42, 14, 54, 26], outline=(200, 190, 170, 255), width=3)
    d.line([(48, 26), (48, 44)], fill=(210, 200, 180, 255), width=2)
    d.polygon([(40, 44), (56, 44), (52, 72), (44, 72)], fill=(230, 220, 200, 255), outline=(160, 145, 120, 255))


def bronze_arrowhead(im, d):
    d.polygon([(48, 16), (68, 58), (48, 50), (28, 58)], fill=(190, 130, 70, 255), outline=(120, 80, 40, 255))
    d.rectangle([44, 50, 52, 82], fill=(140, 100, 60, 255))


def goldsmith_earring(im, d):
    d.ellipse([42, 14, 54, 26], outline=(220, 180, 60, 255), width=3)
    d.line([(48, 26), (48, 40)], fill=(200, 160, 50, 255), width=2)
    d.ellipse([38, 40, 58, 68], fill=(230, 190, 70, 255), outline=(160, 120, 40, 255))
    d.ellipse([44, 46, 52, 54], fill=(250, 230, 140, 255))


def leather_bracer(im, d):
    d.rounded_rectangle([28, 28, 68, 72], radius=6, fill=(120, 75, 45, 255), outline=(70, 45, 25, 255))
    d.rectangle([32, 36, 64, 42], fill=(90, 55, 30, 255))
    d.rectangle([32, 52, 64, 58], fill=(90, 55, 30, 255))
    d.ellipse([44, 44, 52, 52], fill=(160, 120, 60, 255))


def linen_capelet(im, d):
    d.pieslice([18, 28, 78, 88], 200, 340, fill=(230, 220, 200, 255), outline=(160, 150, 130, 255))
    d.ellipse([40, 24, 56, 40], fill=(210, 200, 180, 255), outline=(150, 140, 120, 255))


def buckler_wood(im, d):
    d.ellipse([18, 18, 78, 78], fill=(150, 105, 60, 255), outline=(90, 60, 35, 255))
    d.ellipse([30, 30, 66, 66], outline=(110, 75, 40, 255), width=3)
    d.ellipse([42, 42, 54, 54], fill=(180, 140, 80, 255))


def bone_arrowhead(im, d):
    d.polygon([(48, 14), (66, 56), (48, 48), (30, 56)], fill=(230, 220, 200, 255), outline=(160, 145, 120, 255))
    d.rectangle([44, 48, 52, 84], fill=(200, 185, 160, 255))


def leather_gloves(im, d):
    d.rounded_rectangle([30, 34, 66, 78], radius=8, fill=(130, 85, 50, 255), outline=(80, 50, 30, 255))
    for i, x in enumerate([34, 42, 50, 58]):
        d.rounded_rectangle([x, 18 + (i % 2), x + 8, 40], radius=3, fill=(140, 95, 55, 255), outline=(80, 50, 30, 255))


def linen_cuffs(im, d):
    d.rounded_rectangle([24, 36, 72, 68], radius=4, fill=(225, 215, 195, 255), outline=(150, 140, 120, 255))
    d.rectangle([28, 44, 68, 50], fill=(200, 190, 170, 255))
    d.ellipse([44, 52, 52, 60], fill=(180, 160, 100, 255))


DRAWERS = {
    "bone-ring": bone_ring,
    "glass-bead": glass_bead,
    "distilled-ash": distilled_ash,
    "dried-ration": dried_ration,
    "pale-bait": pale_bait,
    "spice-rub": spice_rub,
    "smoked-minnow": smoked_minnow,
    "ash-dust": ash_dust,
    "bone-chip": bone_chip,
    "dried-ration-scrap": dried_ration_scrap,
    "desiccated-hide": desiccated_hide,
    "antidote": antidote,
    "ashbeam-cane": ashbeam_cane,
    "bone-earring": bone_earring,
    "bronze-arrowhead": bronze_arrowhead,
    "goldsmith-earring": goldsmith_earring,
    "leather-bracer": leather_bracer,
    "linen-capelet": linen_capelet,
    "buckler-wood": buckler_wood,
    "bone-arrowhead": bone_arrowhead,
    "leather-gloves": leather_gloves,
    "linen-cuffs": linen_cuffs,
}


def mark_manifest_done(slugs: set[str]) -> None:
    if not MANIFEST.exists():
        print("manifest missing, skip status update")
        return
    wb = load_workbook(MANIFEST)
    ws = wb["Master Item List"]
    updated = 0
    for row in ws.iter_rows(min_row=2):
        slug_cell = row[1]
        status_cell = row[0]
        slug = str(slug_cell.value or "").strip()
        if slug in slugs and str(status_cell.value or "").strip().lower() != "done":
            status_cell.value = "Done"
            updated += 1
    wb.save(MANIFEST)
    print(f"manifest Status->Done for {updated} rows")


def main() -> None:
    made: set[str] = set()
    for slug, drawer in DRAWERS.items():
        im, d = new_canvas()
        drawer(im, d)
        save(im, f"{slug}.png")
        made.add(slug)
    mark_manifest_done(made)
    print(f"done {len(made)} icons")


if __name__ == "__main__":
    main()
